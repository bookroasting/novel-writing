"""publish 단계 1단계: 원고 → Word 문서 (범용).

사용:
  python3 <스킬>/scripts/generate_docx.py --run <run> --trim 신국판 [--root <프로젝트 폴더>]

표지·판권·목차·폰트·여백·판형은 모두 book-toc.md의 작품 파라미터 블록에서 읽는다.
원고에서는 블록에 정의된 부·장 구간과 `back_matter` 섹션만 옮긴다.
"""
import argparse
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

try:
    from docx import Document
except ImportError:
    raise SystemExit(f"python-docx가 없다. python3 -m pip install -r {Path(__file__).resolve().parents[1] / 'requirements.txt'} 로 설치한 뒤 다시 실행한다")
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from booktoc import require_run, ROOT, DEFAULT_TOC, load_params, heading, latest_run, split_draft, outline, final_draft_name, run_params, split_back, back_headings, ledger_trim, TRIM_CM  # noqa: E402

GRAY = RGBColor(0x88, 0x88, 0x88)


def set_font(run, name, size, bold=False, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for k in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rf.set(qn(k), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


SERIF_DEFAULTS = ["Nanum Myeongjo", "AppleMyungjo", "Batang", "Noto Serif KR"]     # 흔히 깔린 명조·바탕 계열
SANS_DEFAULTS = ["Nanum Gothic", "Apple SD Gothic Neo", "Malgun Gothic", "Noto Sans KR"]   # 흔히 깔린 고딕·돋움 계열
# 글꼴 계열 정보: 대체 글꼴을 고를 때 Word·LibreOffice가 계열(명조/고딕)을 보고 비슷한 것을 찾게 한다
FAMILY = {"roman": ("roman", "02030600000101010101"), "swiss": ("swiss", "020B0600000101010101")}


def font_entry(name, alt, kind):
    fam, panose = FAMILY[kind]
    alt_xml = f'<w:altName w:val="{alt}"/>' if alt else ""
    return (f'<w:font w:name="{name}">{alt_xml}<w:panose1 w:val="{panose}"/><w:charset w:val="81"/>'
            f'<w:family w:val="{fam}"/><w:pitch w:val="variable"/></w:font>')


FONT_DIRS = ["/System/Library/Fonts", "/Library/Fonts", "~/Library/Fonts",            # macOS
             "C:/Windows/Fonts", "~/AppData/Local/Microsoft/Windows/Fonts",           # Windows
             "/usr/share/fonts", "/usr/local/share/fonts", "~/.local/share/fonts", "~/.fonts"]   # Linux


def font_families(path):
    """TTF·OTF·TTC 파일의 name 표에서 글꼴 계열 이름(nameID 1, 16)을 읽는다. fontconfig 없이도 동작한다."""
    import struct
    names = set()
    try:
        data = Path(path).read_bytes()
    except OSError:
        return names
    offs = [0]
    if data[:4] == b"ttcf":
        n = struct.unpack(">I", data[8:12])[0]
        offs = list(struct.unpack(f">{n}I", data[12:12 + 4 * n]))
    for off in offs:
        try:
            num = struct.unpack(">H", data[off + 4:off + 6])[0]
            for k in range(num):
                tag, _, toff, _ = struct.unpack(">4sIII", data[off + 12 + 16 * k: off + 28 + 16 * k])
                if tag != b"name":
                    continue
                _, count, soff = struct.unpack(">HHH", data[toff:toff + 6])
                for r in range(count):
                    pid, eid, lid, nid, ln, o = struct.unpack(">HHHHHH", data[toff + 6 + 12 * r: toff + 18 + 12 * r])
                    if nid not in (1, 16):
                        continue
                    raw = data[toff + soff + o: toff + soff + o + ln]
                    names.add((raw.decode("utf-16-be", "ignore") if pid in (0, 3) else raw.decode("mac_roman", "ignore")).strip())
        except struct.error:
            continue
    return {n for n in names if n}


def installed_fonts():
    """이 컴퓨터에 깔린 글꼴 계열 이름 집합. fc-list가 있으면 그것을, 없으면 글꼴 폴더를 직접 읽는다(macOS·Windows·Linux)."""
    import os
    try:
        out = __import__("subprocess").run(["fc-list", ":", "file", "family"], capture_output=True, text=True, timeout=20).stdout
        if out.strip():
            # macOS가 내려받기만 하고 켜지 않은 글꼴(AssetsV2)은 프로그램이 쓰지 못하므로 뺀다
            return {n.strip() for ln in out.splitlines() if "/AssetsV2/" not in ln.split(":")[0]
                    for n in ln.split(":", 1)[-1].split(",") if n.strip()}
    except Exception:  # noqa: BLE001
        pass
    found = set()
    import glob
    for d in FONT_DIRS:
        for base in map(Path, glob.glob(os.path.expanduser(d))):
            for p in base.rglob("*"):
                if p.suffix.lower() in (".ttf", ".otf", ".ttc"):
                    found |= font_families(p)
    return found or None


def add_font_fallbacks(path, fonts):
    """docx 글꼴 표(fontTable.xml)에 글꼴마다 계열(명조=roman, 고딕=swiss)과 대체 이름(altName) 하나를 적는다.
    fonts: [(글꼴 이름, 대체 후보 목록, 계열)]. 대체 이름은 후보 중 이 컴퓨터에 깔린 첫 글꼴, 없으면 첫 후보다.
    docx에는 글꼴 이름이 하나만 저장되므로, 읽는 쪽에 원래 글꼴이 없을 때 계열이 맞는 글꼴로 바뀌게 하는 장치다."""
    have = installed_fonts()
    with zipfile.ZipFile(path) as z:
        items = {i.filename: z.read(i.filename) for i in z.infolist()}
    ft = items["word/fontTable.xml"].decode("utf-8")
    report = []
    for name, cands, kind in fonts:
        alt = next((c for c in cands if have and c in have), cands[0] if cands else "")
        ft = re.sub(rf'<w:font w:name="{re.escape(name)}"(?:/>|>.*?</w:font>)', "", ft, flags=re.S)
        ft = ft.replace("</w:fonts>", font_entry(name, alt, kind) + "</w:fonts>")
        if have is not None and name not in have:
            report.append(f"{name}이(가) 이 컴퓨터에 없다. 글꼴 표에 계열 정보와 대체 이름 '{alt}'을 적었다. "
                          f"Word는 이 정보로 비슷한 글꼴을 고르지만 LibreOffice·Pages는 자기 대체 표를 써서 손글씨·장식 글꼴이 나올 수 있다. "
                          f"확실한 조판이 필요하면 {name}을 설치하거나 블록 글꼴을 깔린 글꼴로 바꾼다(작가 승인 gate: block-change)")
    items["word/fontTable.xml"] = ft.encode("utf-8")
    tmp = Path(tempfile.mkstemp(suffix=".docx")[1])
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in items.items():
            z.writestr(k, v)
    shutil.move(str(tmp), str(path))
    return report


def centered(doc, text, font, size, bold=False, color=None, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), font, size, bold, color)
    return p


def build(draft_text, p, trim, out):
    pub = p["publish"]
    hc = RGBColor.from_string(pub["heading_color"])
    chapters, _ = split_draft(draft_text, p)

    doc = Document()
    sec = doc.sections[0]
    w, h = TRIM_CM[trim]
    sec.page_width, sec.page_height = Cm(w), Cm(h)
    m = pub["margin_cm"]
    sec.top_margin, sec.bottom_margin = Cm(m["top"]), Cm(m["bottom"])
    sec.left_margin, sec.right_margin = Cm(m["left"]), Cm(m["right"])
    doc.core_properties.title = p["title"]
    doc.core_properties.author = p["pen_name"]

    # 앞부분(표지·판권·목차)은 첫 구역, 본문은 둘째 구역이다. 머리글·쪽번호는 본문 구역에만 넣고 쪽번호는 1부터 센다.
    # 세로 위치는 빈 문단이 아니라 판형 높이에 비례한 문단 앞 간격으로 정한다(Word·LibreOffice에서 같은 자리).
    usable_pt = (h - m["top"] - m["bottom"]) * 28.3465

    # 표지
    centered(doc, p["title"], pub["heading_font"], 36, True, hc, before=usable_pt * 0.28)
    centered(doc, pub["cover_label"], pub["body_font"], 14, color=GRAY, before=24)
    centered(doc, p["pen_name"], pub["body_font"], 12, before=usable_pt * 0.30)

    # 판권 (e북과 같은 블록 값. 비어 있는 값의 줄은 넣지 않는다). 쪽 아래쪽에 둔다
    colophon = [(p["title"], True), (pub["cover_label"], False), (f"지은이 · {p['pen_name']}", False)]
    if pub.get("pub_date"):
        colophon.append((f"{pub['pub_date']} 발행", False))
    if pub.get("publisher"):
        colophon.append((f"발행 · {pub['publisher']}", False))
    if pub.get("isbn"):
        colophon.append((f"ISBN {pub['isbn']}", False))
    if pub.get("copyright_year"):
        colophon.append((f"© {pub['copyright_year']} {pub.get('copyright_holder') or p['pen_name']}", False))
        colophon.append(("All rights reserved.", False))
    for k, (text, bold) in enumerate(colophon):
        cp = centered(doc, text, pub["body_font"], 9, bold, GRAY, before=(usable_pt * 0.55 if k == 0 else 12 if k == 2 else 0), after=2)
        if k == 0:
            cp.paragraph_format.page_break_before = True

    # 목차
    centered(doc, "목차", pub["heading_font"], 20, True, hc, after=24).paragraph_format.page_break_before = True
    backs = split_back(draft_text, p)
    for kind, h in outline(p) + [("back", bh) for bh, _ in backs]:
        centered(doc, h[2:], pub["heading_font"] if kind == "part" else pub["body_font"], pub["body_size_pt"],
                 bold=(kind == "part"), after=8)

    # 본문 (장편이면 부 제목 페이지가 해당 장들 앞에 온다)
    by_heading = {heading(c): c for c in p["chapters"]}
    back_text = dict(backs)
    def new_section(restart=False):
        """새 쪽에서 시작하는 구역을 연다. add_section이 만든 빈 문단은 지우고 구역 정보를 앞 문단으로 옮긴다."""
        sec_ = doc.add_section(WD_SECTION.NEW_PAGE)
        holder = [q for q in doc.paragraphs if q._p.pPr is not None and q._p.pPr.find(qn("w:sectPr")) is not None][-1]
        prev = holder._p.getprevious()
        if not holder.text.strip() and prev is not None and prev.tag == qn("w:p"):
            prev.get_or_add_pPr().append(holder._p.pPr.find(qn("w:sectPr")))
            holder._p.getparent().remove(holder._p)
        for old in sec_._sectPr.findall(qn("w:pgNumType")):   # 앞 구역에서 복사된 번호 다시 세기를 지운다
            sec_._sectPr.remove(old)
        if restart:
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:start"), "1")
            sec_._sectPr.append(pg)
        return sec_

    def header_footer(sec_, on):
        """구역의 머리글·바닥글. on이면 블록 header·footer대로, 아니면 비운다(표지·판권·목차·부 제목 쪽)."""
        for part in (sec_.header, sec_.footer):
            part.is_linked_to_previous = False
            for para in part.paragraphs:
                for r_ in list(para.runs):
                    r_._element.getparent().remove(r_._element)
        if not on:
            return
        if pub.get("header") == "title":
            hp_ = sec_.header.paragraphs[0]
            hp_.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(hp_.add_run(p["title"]), pub["body_font"], 9, color=GRAY)
        if pub.get("footer") == "page_number":
            fp = sec_.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = fp.add_run()
            for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
                el = OxmlElement(tag)
                if attr:
                    el.set(qn("w:fldCharType"), attr)
                else:
                    el.text = "PAGE"
                run._element.append(el)
            set_font(run, pub["body_font"], 9, color=GRAY)

    # 본문 구역들: 쪽번호는 본문 첫 쪽이 1. 부 제목 쪽은 따로 구역을 열어 머리글·쪽번호를 비운다(번호는 이어서 센다)
    sec_kinds = ["front"]   # 만든 순서대로 구역 종류를 적는다(python-docx 구역 객체는 비교할 수 없어 순서로 맞춘다)
    new_section(restart=True)
    sec_kinds.append(None)
    cur_kind = None
    first_page = True
    for kind, h in outline(p) + [("back", bh) for bh, _ in backs]:
        want = "plain" if kind == "part" else "body"
        if cur_kind is None:
            cur_kind = want
        elif want != cur_kind or kind == "part":
            new_section()
            sec_kinds.append(None)
            cur_kind = want
            first_page = True   # 구역 나눔이 새 쪽을 연다
        sec_kinds[-1] = want
        new_page = not first_page   # 쪽 나눔은 빈 문단이 아니라 제목 문단의 '앞에서 쪽 나눔'으로 한다
        first_page = False
        if kind == "part":
            centered(doc, h[2:], pub["heading_font"], pub["heading_size_pt"] + 4, True, hc,
                     before=usable_pt * 0.35).paragraph_format.page_break_before = new_page
            continue
        if kind == "back":
            c, text = None, back_text[h]
        else:
            c = by_heading[h]
            text = chapters.get(c["no"], "")
        hp = doc.add_paragraph()
        hp.paragraph_format.page_break_before = new_page
        hp.paragraph_format.space_before = Pt(60)
        hp.paragraph_format.space_after = Pt(36)
        set_font(hp.add_run(h[2:]), pub["heading_font"], pub["heading_size_pt"], True, hc)
        # 문단 사이 빈 줄은 빈 문단으로 옮기지 않는다(문단 간격은 space_after가 맡고, e북과 같은 문단 구성이 된다).
        # 빈 줄이 두 줄 이상이면 장면 전환으로 보고 다음 문단 앞에 간격을 둔다.
        blanks = 0
        for line in text.strip().split("\n"):
            if not line.strip():
                blanks += 1
                continue
            bp = doc.add_paragraph()
            f = bp.paragraph_format
            if blanks >= 2:
                f.space_before = Pt(pub["body_size_pt"] * pub["line_spacing"] * 1.5)
            blanks = 0
            f.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            f.line_spacing = pub["line_spacing"]
            f.space_after = Pt(4)
            f.first_line_indent = Pt(pub["body_size_pt"] * pub["first_line_indent_chars"])
            set_font(bp.add_run(line.strip()), pub["body_font"], pub["body_size_pt"])

    for sec_, kind_ in zip(doc.sections, sec_kinds):
        if kind_ != "front":
            header_footer(sec_, kind_ == "body")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    body_fb = list(dict.fromkeys((pub.get("body_font_fallback") or []) + SERIF_DEFAULTS))
    head_fb = list(dict.fromkeys((pub.get("heading_font_fallback") or []) + SANS_DEFAULTS))
    for msg in add_font_fallbacks(out, [(pub["body_font"], body_fb, "roman"), (pub["heading_font"], head_fb, "swiss")]):
        print(f"WARN 글꼴: {msg}")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run")
    ap.add_argument("--draft")
    ap.add_argument("--out")
    ap.add_argument("--toc", default=str(DEFAULT_TOC))
    ap.add_argument("--root", help="프로젝트 폴더 (기본: book-toc.md를 위로 찾아감)")
    ap.add_argument("--trim")
    a = ap.parse_args()
    run = require_run(a.run or latest_run())
    p = run_params(run) if a.toc == str(DEFAULT_TOC) else load_params(a.toc)
    draft = Path(a.draft) if a.draft else ROOT / "02_draft" / run / final_draft_name(p)
    out = Path(a.out) if a.out else ROOT / "03_output" / run / "final.docx"
    decided = ledger_trim(run) if a.run or not a.draft else None
    if a.trim and decided and a.trim != decided:
        raise SystemExit(f"원장의 판형 결정은 '{decided}'이다. 바꾸려면 사용자에게 다시 묻고 gate: trim 줄을 새로 남긴다")
    trim = a.trim or decided
    if not trim:
        raise SystemExit(f"판형이 정해지지 않았다. 블록 trim_options {p['publish']['trim_options']} 중에서 사용자에게 묻고 "
                         f"원장에 gate: trim → <판형> \"<원문>\"을 남긴 뒤 다시 실행한다")
    if trim not in TRIM_CM:
        raise SystemExit(f"알 수 없는 판형: {trim} (가능: {list(TRIM_CM)})")
    build(draft.read_text(encoding="utf-8"), p, trim, out)
    print(f"OK {out}  ({p['title']} / {p['pen_name']} / {trim})")


if __name__ == "__main__":
    main()
