"""범용성 회귀 시험.

사용: python3 tools/tests/selftest.py   (또는 bash tools/tests/selftest.sh)

T1 다른 장르 픽스처(가족 드라마 40매, 3장)의 정상 원고가 검증기를 통과한다
T2 같은 픽스처의 결함 원고를 검증기가 결함별로 잡는다
T3 범용 docx 생성기가 픽스처로 표지·목차·마지막 문단이 맞는 docx를 만든다
T4 현재 프로젝트가 린터를 통과한다
T5 리팩터 이전 스킬·카드(_archive)를 넣은 사본에서 린터가 L1·L3·L4를 잡는다
T6 README 절차대로 book-toc·storyline 두 파일만 바꾼 사본에서 린터 FAIL 0 (지난 run은 스냅샷으로 검사)
T7 장편(part) 모드: 부 헤딩이 있는 원고 통과, 부 헤딩 누락 탐지, docx 목차에 부 제목
T8 실행 감사: 합평 머리 줄 누락, 점수 동결 위반, 바뀐 자리 누락, 승인 없는 waiver를 잡고 정상 run은 통과
T9 빈칸이 남은 템플릿 블록을 린터가 L2로 잡는다
T10 범용 e북 생성기: 다른 장르를 plain 디자인으로 만들고 검증 통과, 로봇 그림·이전 작품 문자열 없음, back_matter 포함
T11 장편 부 파일을 validate_draft.py CLI에 여러 개로 넘겨 통과 (스킬에 적힌 명령 그대로)
T12 우회 차단: 없는 입력 파일, rejected 응답, 승인 없는 run 종료를 린터가 FAIL로 잡는다
T13 new_run.py rework: 끝난 run을 물려받아 /review부터 시작할 run을 만들고 린터 FAIL 0
T14 통과 판정을 블록 통과선과 비교 (미달이면 FAIL, 승인된 open_red면 WARN)
T15 e북만 출판하는 run도 출판 감사를 받는다
T16 떨어진 라운드끼리도 같은 입력이면 동결을 강제한다
T17 e북: 제사가 없으면 본문 시작 쪽을 계산하고, 빈 판권 줄을 넣지 않는다
T18 설치: 스킬 폴더가 그대로 복사되고, 서브 에이전트의 {{SKILL_DIR}}이 설치 경로로 바뀌며, 설치본 스크립트가 프로젝트를 찾는다
T19 블록 스키마: 범위 밖 값, 빈 목록 항목, 없는 장 번호, e북 하위 필드 형식, JSON 오류 안내를 줄 번호와 함께 잡는다
T20 waiver 1:1: 승인 줄 하나가 앞선 waiver 둘을 함께 승인하지 않는다
T21 게이트 순서: 본문 합평 미통과인데 최종본이 있으면 FAIL, review-9 reopen은 본문 통과를 무르지 않는다
T22 작가 산출물 제거: MANIFEST로 옮겨 없어진 03_output 파일에 remove 승인이 없으면 FAIL
T23 rework 우회 차단: 스토리 게이트가 닫히지 않은 run은 new_run이 물려받지 않는다
T24 storyline 양식 빈칸이 남으면 FAIL
T25 단계 건너뛰기: 원장에 publish done만 있고 앞 단계·판형 게이트가 없으면 FAIL
T26 gate_status: rework 안의 reopen 이후 재합평, 장편의 빠진 부, 본문 reopen 뒤 새 라운드 없음, 다른 kind의 gate_log waiver
T27 rework가 서평 결정을 승인 원문 그대로, 판형 결정과 함께 물려주고, back_matter 섹션을 빼면 검증기가 FAIL
T28 잠금: 잠금 뒤 back_matter·formats가 승인 없이 줄거나 스냅샷이 기록 없이 바뀌면 FAIL, 승인이 있으면 통과, 출판 뒤 잠금 누락 FAIL
T29 첫 작품 시작: init_project.py가 답 한 묶음으로 두 파일을 만들고, 블록은 린터 스키마를 통과하며, 있는 파일은 덮어쓰지 않는다
T30 출판 감사: 원장 판형과 다른 docx, 승인 없는 미확인 제사를 FAIL로 잡고, docx 생성기는 판형 결정 없이 멈춘다
T31 e북: 오프라인 대체 화면이 들어 있고, 디자인 전용 조각은 값이 채워져 슬롯으로 옮겨지며 결과물에 template이 남지 않는다
T32 gate_status: 스토리용 gate_log waiver는 본문 게이트를 닫지 않는다
T33 잠금 우회 차단: 다시 잠그기 거부, 잠금 구간마다 비교, rejected·note는 승인이 아님
T34 waiver 짝: 승인 줄은 바로 위 waiver 하나만 승인한다
T35 출판 뒤 수정: republish.py는 승인 없이 멈추고, 승인이 있으면 다시 만들고 메타 분량을 맞춰 잠근다
T36 rework가 출판 뒤 작가 수정 기록을 넘긴다
T37 docx 조판: 앞부분 구역에는 머리글·쪽번호가 없고 본문 구역 쪽번호는 1부터, 빈 문단 대신 간격, 글꼴 표에 대체 글꼴
T38 e북 로더: 200 응답이어도 라이브러리가 없으면 다음 후보·대체 화면으로, 글꼴 대기는 1.5초 상한
T39 잠금 키 범위: 무관한 승인(post-audit)·note로는 다른 키를 못 바꾸고, 키를 적은 block-change 승인은 통과
T40 잠금 줄·잠금 파일 삭제를 잡는다
T41 출판 뒤 잠긴 최종 원고를 승인 없이 고치면(글자 수가 같아도) 잡는다, 작가의 유지 결정이 있으면 republish.py가 멈춘다
T42 잠금 파일 위조·산출물 직접 재생성·응답 원문 속 낱말을 키로 쓰기를 잡고, 출판 뒤 수정은 owner만 마치며 진행 중에는 두 번째 수정과 rework가 열리지 않는다
T43 승인 전에 이미 고친 원고는 --prepare가 '수정 전 보관본'으로 받지 않는다
T44 fc-list 없이도 글꼴 폴더에서 글꼴 계열 이름을 읽는다
T45 끊긴 출판 뒤 수정: --status가 알리고, 승인 없이 이어받기·버리기가 멈추며, 버리기는 원고를 되돌려 다음 수정을 열고, 이어받기는 새 owner로 마친다
T46 출판 뒤 수정 우회: owner 없는 prepare 줄 덧붙이기, 손으로 쓴 done 줄과 재잠금, prepare 없는 수정을 잡는다
T47 승인 범위: 겹치는 끝 이름(lines)은 키가 아니고, publish-info는 빈칸 채우기만, epigraph-source는 출처 표시만 덮는다
T48 출판 시점에 metadata 분량 줄 형식을 잡고, 없는 run에는 스크립트마다 있는 run 목록으로 답한다
T49 장편 docx: 부 제목 쪽에는 머리글·쪽번호가 없고 장 쪽에는 있으며, 쪽번호는 본문 첫 구역에서만 1로 시작한다
T50 출판된 docx가 지금 생성기 결과와 다르면 린터가 알린다
T51 버리기: 끊긴 세션이 산출물까지 바꿨으면 산출물도 되돌리고, --status는 요청 원문·변경 여부를 보여 준다
"""
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
if (REPO / "SKILL.md").exists():   # 스킬 저장소(루트가 스킬): 예시 작품 폴더를 프로젝트로 쓴다
    SKILL, ROOT = REPO, REPO / "examples" / "kimjang-day"
else:                              # 작품 프로젝트 안의 스킬(novel-writing/ 하위 폴더)
    SKILL, ROOT = REPO / "novel-writing", REPO
SCRIPTS = SKILL / "scripts"
sys.path.insert(0, str(SCRIPTS))
from booktoc import load_params, lock_line  # noqa: E402
from validate_draft import validate  # noqa: E402
from generate_docx import build  # noqa: E402
from lint_project import lint  # noqa: E402
from generate_ebook import build as build_ebook  # noqa: E402
from validate_draft import check_ebook  # noqa: E402

FIX = HERE / "fixtures" / "family" / "book-toc.md"
results = []


def check(name, ok, detail=""):
    ok = bool(ok)   # 빈 목록 같은 거짓 값도 FAIL로 센다
    results.append(ok)
    print(f"{'PASS' if ok else 'FAIL'} {name}{'  ' + detail if detail and not ok else ''}")


def make_draft(p, extra=""):
    """블록 목표 분량에 맞는 합성 원고. 문장 내용은 시험용이다."""
    filler = ["엄마가 소매를 걷었다.", "배추가 마당에 쌓였다.", "윤재가 물을 길어 왔다.", "나는 칼을 닦았다.",
              "바람이 조금 불었다.", "엄마가 고무장갑을 꼈다.", "소금 자루가 무거웠다.", "우리는 말없이 일했다."]
    out = []
    for c in p["chapters"]:
        target = c["pages"] * p["length"]["chars_per_page"]
        out.append(f"# {c['no']}장. {c['title']}\n")
        body, i = [], 0
        while sum(len(x) + 1 for x in body) < target - 40:
            body.append(filler[i % len(filler)])
            i += 1
        lines = [" ".join(body[j:j + 4]) for j in range(0, len(body), 4)]
        if c["no"] == p["chapters"][-1]["no"]:
            lines.append(p["last_line"])
        out.append("\n\n".join(lines) + "\n")
    return "\n".join(out) + extra


def project_copy(td):
    """린터에 필요한 최소 사본: 작품 파일, 진입 문서, 00_user_input, run 폴더, 보관 기록. 스킬은 원본을 그대로 쓴다."""
    t = Path(td)
    for d in ("00_user_input", "01_test", "02_draft", "03_output", "designs"):
        if (ROOT / d).exists():
            shutil.copytree(ROOT / d, t / d, ignore=shutil.ignore_patterns("__pycache__"))
    for f in ("AGENTS.md", "CLAUDE.md", "book-toc.md"):
        if (ROOT / f).exists():
            shutil.copy(ROOT / f, t / f)
    (t / "_archive").mkdir()
    for f in ("MANIFEST.tsv", "retired_works.json"):
        if (ROOT / "_archive" / f).exists():
            shutil.copy(ROOT / "_archive" / f, t / "_archive" / f)
    # 실제 작품에 다른 세션의 출판 뒤 수정이 진행 중이면, 사본에서만 그 수정을 닫아 둔다(rework·prepare 시험이 막히지 않게)
    # 절차 그대로: 그 수정의 보관본을 사본에 가져와 작가 승인 줄을 쓰고 --abandon으로 되돌린다
    from booktoc import open_post_audit
    for d in (t / "02_draft").glob("*_*"):
        st = open_post_audit(d.name, t) if (d / "00_RUN_STATUS.md").exists() else None
        if not st:
            continue
        led = d / "00_RUN_STATUS.md"
        arc = re.match(r"^note:\s*post-audit-prepare\s+(\S+)", led.read_text(encoding="utf-8").split("\n")[st[0]])
        if arc and (ROOT / arc.group(1)).exists():
            shutil.copytree(ROOT / arc.group(1), t / arc.group(1))
        led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + '\ngate: post-audit-abandon → approved (selftest) "사본에서만 닫는다"\n', encoding="utf-8")
        subprocess.run([sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", d.name, "--root", str(t), "--abandon"],
                       capture_output=True, text=True, cwd=t)
    return t


def write_block(path, params):
    text = path.read_text(encoding="utf-8")
    new = "```json\n" + json.dumps(params, ensure_ascii=False, indent=2) + "\n```"
    path.write_text(re.sub(r"```json\s*\n.*?\n```", lambda _: new, text, count=1, flags=re.S), encoding="utf-8")


def ensemble(inp, score, body="", verdict="재수정", reds=0, name=None, digest=None):
    h = digest or hashlib.sha256(inp.read_bytes()).hexdigest()
    return f"입력: {name or inp.name}\n입력 sha256: {h}\n판정 점수: {score}\n남은 🔴: {reds}\n판정: {verdict}\n\n{body}\n"


def lint_codes(root):
    r = lint(root)
    return r, " | ".join(r.fails)


def main():
    p = load_params(FIX)

    # T1
    good = make_draft(p)
    fails, warns, info = validate(good, p)
    check("T1 다른 장르 정상 원고 통과", not fails, str(fails))

    # T2
    bad = good.replace("나는 칼을 닦았다.", "윤서는 칼을 닦았다 — 천천히.", 1)
    bad = bad.replace("# 2장. 소금\n", "", 1)
    bad += "\n# 작가의 말\n\n고마운 분들께.\n"
    fails, _, _ = validate(bad, p)
    joined = " | ".join(fails)
    for label, key in [("em dash", "em dash"), ("1인칭 시점", "1인칭 시점 위반"), ("장 누락", "장 헤딩 누락"),
                       ("장 밖 섹션", "허용되지 않은 장 밖 섹션")]:
        check(f"T2 결함 탐지: {label}", key in joined, joined)

    # T3
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "final.docx"
        build(good, p, "신국판", out)
        from validate_draft import check_docx
        last = [x for x in good.strip().split("\n") if x.strip()][-1]
        dfails = check_docx(out, p, last)
        check("T3 범용 docx 생성기 (표지·목차·마지막 문단)", not dfails, str(dfails))

    # T4
    # 실제 작품의 지금 상태는 도구 회귀와 따로 본다. 진행 중인 작가 작업이 도구 시험을 깨뜨리지 않게, 이후 시험은 이 기준선에 없던 FAIL만 센다
    r = lint(ROOT)
    BASELINE = set(r.fails)
    # 출판까지 마친 가장 최근 run (rework·republish 시험의 재료). 없으면 그 시험은 건너뛴다
    done = [d.name for d in sorted((ROOT / "02_draft").glob("*_*")) if (d / "00_RUN_STATUS.md").exists()
            and re.search(r"^stage:\s*publish done", (d / "00_RUN_STATUS.md").read_text(encoding="utf-8"), re.M)]
    LIVE = done[-1] if done else None
    print(f"INFO T4 현재 프로젝트 린터: FAIL {len(r.fails)}건 (도구 회귀가 아니라 작품 상태다. lint_project.py로 따로 본다)"
          + ("".join(f"\n     {x}" for x in r.fails[:3]) if r.fails else ""))

    # T5
    snap = ROOT / "_archive" / "20261003_pre-skill-refactor"
    if not snap.exists():
        print("SKIP T5 리팩터 이전 스킬 보관본이 이 저장소에 없다(스킬만 받은 저장소)")
    else:
        with tempfile.TemporaryDirectory() as td:
            t = project_copy(td)
            sk = t / "skill"   # 리팩터 이전 문서를 담은 스킬 사본
            shutil.copytree(SKILL, sk, ignore=shutil.ignore_patterns("__pycache__"))
            shutil.copy(snap / ".claude" / "reviewers.md", sk / "references" / "reviewers.md")
            shutil.copy(snap / ".claude" / "skills" / "write.md", sk / "references" / "stage-write.md")
            (t / ".claude" / "skills").mkdir(parents=True)
            shutil.copy(snap / ".claude" / "skills" / "review.md", t / ".claude" / "skills" / "review.md")
            r = lint(t, sk)
            codes = {f.split()[0] for f in r.fails}
            for c in ("L1", "L3", "L4"):
                check(f"T5 리팩터 이전 문서에서 {c} 탐지", c in codes, str(sorted(codes)))

    # T6
    with tempfile.TemporaryDirectory() as td:
        t6 = project_copy(td)
        shutil.copy(FIX, t6 / "book-toc.md")
        shutil.copy(HERE / "fixtures" / "family" / "storyline.md", t6 / "00_user_input" / "storyline.md")
        r, j = lint_codes(t6)
        new6 = [x for x in r.fails if x not in BASELINE]
        check("T6 두 파일만 교체한 사본에서 새 린터 FAIL 0", not new6, " | ".join(new6)[:300])

    # T7
    pp = json.loads(json.dumps(p))
    pp["structure_mode"] = "part"
    pp["parts"] = [{"no": 1, "title": "가을", "chapters": [1, 2]}, {"no": 2, "title": "겨울", "chapters": [3]}]
    txt = make_draft(pp).replace("# 1장. 배추", "# 1부. 가을\n\n# 1장. 배추").replace("# 3장. 항아리", "# 2부. 겨울\n\n# 3장. 항아리")
    f7, _, _ = validate(txt, pp)
    check("T7 장편 원고 통과", not f7, str(f7))
    f7b, _, _ = validate(txt.replace("# 2부. 겨울\n", ""), pp)
    check("T7 부 헤딩 누락 탐지", any("헤딩" in x for x in f7b), str(f7b))
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "p.docx"
        build(txt, pp, "B5", out)
        from validate_draft import check_docx
        d7 = check_docx(out, pp, pp["last_line"])
        check("T7 장편 docx 목차에 부 제목 포함", not d7, str(d7))

    # T8
    with tempfile.TemporaryDirectory() as td:
        t8 = project_copy(td)
        run = "20990101_01"
        dt, dd = t8 / "01_test" / run, t8 / "02_draft" / run
        dt.mkdir(parents=True)
        dd.mkdir(parents=True)
        shutil.copy(t8 / "book-toc.md", dt / "book-toc.snapshot.md")
        led = dd / "00_RUN_STATUS.md"
        led.write_text("status: in-progress\nstage: write done 2099-01-01\n", encoding="utf-8")
        (dd / "05_draft-v2.md").write_text("v2", encoding="utf-8")
        (dd / "07_draft-v3.md").write_text("v2", encoding="utf-8")
        (dd / "06_ensemble-1.md").write_text(ensemble(dd / "05_draft-v2.md", 8.5), encoding="utf-8")
        (dd / "08_ensemble-2.md").write_text(ensemble(dd / "07_draft-v3.md", 9.1, "## 바뀐 자리\n없음"), encoding="utf-8")
        _, j = lint_codes(t8)
        check("T8 동결 위반 탐지 (같은 입력에 점수 상승)", "[review_freeze]" in j, j[:200])
        (dd / "08_ensemble-2.md").write_text(ensemble(dd / "07_draft-v3.md", 8.5, "점수 동결: 입력 해시가 06 라운드와 같다", verdict="동결"), encoding="utf-8")
        _, j = lint_codes(t8)
        check("T8 정상 동결은 통과", "review_freeze" not in j, j[:200])
        (dd / "07_draft-v3.md").write_text("v3 changed", encoding="utf-8")
        (dd / "08_ensemble-2.md").write_text(ensemble(dd / "07_draft-v3.md", 9.0), encoding="utf-8")
        _, j = lint_codes(t8)
        check("T8 바뀐 자리 누락 탐지", "[change_log]" in j, j[:200])
        (dd / "08_ensemble-2.md").write_text("헤더 없는 합평", encoding="utf-8")
        _, j = lint_codes(t8)
        check("T8 머리 줄 누락 탐지", "[ensemble_header]" in j, j[:200])
        led.write_text(led.read_text(encoding="utf-8") + "waiver: ensemble_header 시험\n", encoding="utf-8")
        _, j = lint_codes(t8)
        check("T8 승인 없는 waiver는 FAIL 유지", "사용자 승인 줄이 없다" in j, j[:200])
        led.write_text(led.read_text(encoding="utf-8") + 'gate: waiver-ensemble_header → approved (2099-01-01) "승인"\n', encoding="utf-8")
        r, j = lint_codes(t8)
        check("T8 승인된 waiver는 WARN으로", "ensemble_header" not in j and any("ensemble_header waived" in w for w in r.warns), j[:200])

    # T9
    with tempfile.TemporaryDirectory() as td:
        t9 = project_copy(td)
        shutil.copy(SKILL / "templates" / "book-toc.template.md", t9 / "book-toc.md")
        _, j = lint_codes(t9)
        check("T9 빈칸 템플릿을 L2로 탐지 (키 경로·줄 번호)", "블록 `title` 빈칸" in j and "서술 섹션 빈칸" in j and "book-toc.md:" in j, j[:200])

    # T10
    pe = json.loads(json.dumps(p))
    pe["back_matter"] = ["작가의 말"]
    txt10 = make_draft(pe) + "\n# 작가의 말\n\n고마운 분들께 드린다.\n"
    f10, _, _ = validate(txt10, pe)
    check("T10 back_matter 섹션이 있는 원고 통과", not f10, str(f10))
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "ebook.html"
        build_ebook(txt10, pe, out)
        h = out.read_text(encoding="utf-8")
        ef, _ = check_ebook(out, pe, "고마운 분들께 드린다.", [x for _, x in __import__("booktoc").outline(pe)] + ["# 작가의 말"])
        check("T10 plain e북 검증 통과", not ef, str(ef))
        from booktoc import past_work_tokens, work_tokens   # 이 저장소의 지난 작품 고유 낱말(run 스냅샷에서). 없으면 현재 블록 낱말
        past = set(past_work_tokens(ROOT)) or (set(work_tokens(load_params(ROOT / "book-toc.md"))) if (ROOT / "book-toc.md").exists() else set())
        leaks = [w for w in sorted(past) if w in h and w not in set(work_tokens(pe))]
        check("T10 plain e북에 이전 작품 문자열 없음", not leaks, str(leaks))

    # T11
    with tempfile.TemporaryDirectory() as td:
        d = Path(td)
        parts_txt = txt.split("# 2부. 겨울")
        (d / "03_draft-v1-part1.md").write_text(parts_txt[0], encoding="utf-8")
        (d / "03_draft-v1-part2.md").write_text("# 2부. 겨울" + parts_txt[1], encoding="utf-8")
        toc = d / "book-toc.md"
        shutil.copy(FIX, toc)
        write_block(toc, pp)
        files = sorted(str(x) for x in d.glob("03_draft-v1-part*.md"))
        cp = subprocess.run([sys.executable, str(SCRIPTS / "validate_draft.py"), "--toc", str(toc), "--draft", *files],
                            capture_output=True, text=True)
        check("T11 장편 부 파일 CLI 검증 통과", cp.returncode == 0, cp.stdout[-300:])

    # T12
    with tempfile.TemporaryDirectory() as td:
        t12 = project_copy(td)
        run = "20990202_01"
        dt, dd = t12 / "01_test" / run, t12 / "02_draft" / run
        dt.mkdir(parents=True)
        dd.mkdir(parents=True)
        shutil.copy(t12 / "book-toc.md", dt / "book-toc.snapshot.md")
        led = dd / "00_RUN_STATUS.md"
        led.write_text("status: in-progress\n", encoding="utf-8")
        (dd / "06_ensemble-1.md").write_text(ensemble(None, 9.8, name="ghost.md", digest="0" * 64), encoding="utf-8")
        _, j = lint_codes(t12)
        check("T12 없는 입력 파일 탐지", "[ensemble_integrity]" in j and "없다" in j, j[:200])
        led.write_text('status: in-progress\nwaiver: ensemble_integrity 시험\ngate: waiver-ensemble_integrity → rejected "안 돼"\n', encoding="utf-8")
        _, j = lint_codes(t12)
        check("T12 rejected 응답은 승인이 아님", "승인 줄" in j, j[:200])
        led.write_text("status: in-progress\nclosed: superseded 2099-02-02\nwaiver: ensemble_integrity 시험\n", encoding="utf-8")
        _, j = lint_codes(t12)
        check("T12 승인 없는 run 종료는 waiver를 인정하지 않음", "승인 줄" in j, j[:200])

    # T13
    if LIVE:  # 출판까지 마친 실제 run이 있어야 하는 시험
        with tempfile.TemporaryDirectory() as td:
            t13 = project_copy(td)
            from new_run import open_run
            run, nxt = open_run(t13, "rework", "다시 합평하자", LIVE, "review", "20991231")
            led = (t13 / "02_draft" / run / "00_RUN_STATUS.md").read_text(encoding="utf-8")
            ok = (nxt.startswith("review 단계") and "stage: write done" in led and f"inherit: {LIVE}" in led
                  and (t13 / "02_draft" / run / "03_draft-v1.md").read_text(encoding="utf-8")
                  == (t13 / "02_draft" / LIVE / "09_draft-final.md").read_text(encoding="utf-8"))
            check("T13 rework run 생성 (원장 상속, 09 → 03)", ok, led[:200])
            r, j = lint_codes(t13)
            new13 = [x for x in r.fails if x not in BASELINE]
            check("T13 rework run을 연 직후 새 린터 FAIL 0", not new13, " | ".join(new13)[:300])

        # T14, T16
        with tempfile.TemporaryDirectory() as td:
            t14 = project_copy(td)
            run = "20990303_01"
            dt, dd = t14 / "01_test" / run, t14 / "02_draft" / run
            dt.mkdir(parents=True)
            dd.mkdir(parents=True)
            shutil.copy(t14 / "book-toc.md", dt / "book-toc.snapshot.md")
            led = dd / "00_RUN_STATUS.md"
            led.write_text("status: in-progress\n", encoding="utf-8")
            (dd / "05_draft-v2.md").write_text("v2", encoding="utf-8")
            (dd / "06_ensemble-1.md").write_text(ensemble(dd / "05_draft-v2.md", 6.1, verdict="통과", reds=4), encoding="utf-8")
            _, j = lint_codes(t14)
            check("T14 기준 미달 통과 판정 탐지", "[pass_threshold]" in j, j[:200])
            led.write_text('status: in-progress\nwaiver: open_red 06_ensemble-1.md 시험\ngate: waiver-open_red → approved (2099-03-03) "진행"\n', encoding="utf-8")
            r, j = lint_codes(t14)
            check("T14 승인된 open_red면 FAIL 아님", "pass_threshold" not in j, j[:200])
            led.write_text("status: in-progress\n", encoding="utf-8")
            (dd / "06_ensemble-1.md").write_text(ensemble(dd / "05_draft-v2.md", 8.0), encoding="utf-8")
            (dd / "07_draft-v3.md").write_text("v3", encoding="utf-8")
            (dd / "08_ensemble-2.md").write_text(ensemble(dd / "07_draft-v3.md", 8.5, "## 바뀐 자리\n1장"), encoding="utf-8")
            (dd / "07_draft-v4.md").write_text("v2", encoding="utf-8")
            (dd / "08_ensemble-3.md").write_text(ensemble(dd / "07_draft-v4.md", 9.9, "## 바뀐 자리\n없음"), encoding="utf-8")
            _, j = lint_codes(t14)
            check("T16 떨어진 라운드의 같은 입력 탐지", "[review_freeze]" in j and "06_ensemble-1.md" in j, j[:300])

    else:
        print("SKIP T13 출판까지 마친 run이 이 저장소에 없다")
    # T15
    with tempfile.TemporaryDirectory() as td:
        t15 = project_copy(td)
        run = "20990404_01"
        dt, dd, do = (t15 / s / run for s in ("01_test", "02_draft", "03_output"))
        for x in (dt, dd, do):
            x.mkdir(parents=True)
        toc = dt / "book-toc.snapshot.md"
        shutil.copy(FIX, toc)
        pf = json.loads(json.dumps(p))
        pf["publish"]["formats"] = ["ebook"]
        write_block(toc, pf)
        bad = make_draft(pf).replace("나는 칼을 닦았다.", "나는 칼을 닦았다 — 천천히.", 1)
        (dd / "09_draft-final.md").write_text(bad, encoding="utf-8")
        build_ebook(bad, pf, do / "ebook.html")
        (do / "metadata.md").write_text("# m\n", encoding="utf-8")
        (do / "validate_report.txt").write_text("x", encoding="utf-8")
        (dd / "00_RUN_STATUS.md").write_text("status: published\nstage: publish done 2099-04-04\n", encoding="utf-8")
        _, j = lint_codes(t15)
        check("T15 e북 전용 run도 출판 감사 (em dash, metadata 섹션)", "[publish_gate]" in j and "em dash" in j and "[metadata]" in j, j[:300])

    # T17
    with tempfile.TemporaryDirectory() as td:
        pe2 = json.loads(json.dumps(p))
        pe2["publish"]["pub_date"] = ""
        pe2["publish"]["copyright_year"] = None
        out = Path(td) / "e.html"
        build_ebook(make_draft(pe2), pe2, out)
        h = out.read_text(encoding="utf-8")
        ok = "src.children.length + tocCount + 1" in h and "bodyStartIdx = 6" not in h and "© None" not in h and "<div> 발행</div>" not in h and "if (false)" in h
        check("T17 제사 없음·빈 판권 처리", ok, "")

    # T18
    sys.path.insert(0, str(REPO / "tools"))
    import install_skill
    with tempfile.TemporaryDirectory() as td:
        d = install_skill.copy_skill(Path(td) / "skills")
        ags = install_skill.copy_agents(Path(td) / "agents", d)
        install_skill.fill_placeholders(d, str(d))
        same = sorted(p.relative_to(d).as_posix() for p in d.rglob("*") if p.is_file()) == \
            sorted(p.relative_to(SKILL).as_posix() for p in SKILL.rglob("*") if p.is_file() and "__pycache__" not in p.parts and p.name != ".DS_Store"
                   and p.relative_to(SKILL).parts[0] in install_skill.SKILL_ITEMS)   # 설치에 들어가는 것만 비교한다
        placeholders = any("{{SKILL_DIR}}" in a.read_text(encoding="utf-8") for a in ags)
        check("T18 설치본 파일 구성이 원본과 같다", same, str(d))
        check("T18 서브 에이전트 5개, 스킬 경로 치환", len(ags) == 5 and not placeholders and all(str(d) in a.read_text(encoding="utf-8") for a in ags), "")
        if LIVE:
            cmd = [sys.executable, str(d / "scripts" / "validate_draft.py"), "--run", LIVE, "--root", str(ROOT)]
        else:   # 출판한 run이 없는 저장소: 픽스처 원고를 설치본 검증기로 검사한다
            fx = Path(tempfile.mkdtemp()) / "draft.md"
            fx.write_text(make_draft(load_params(FIX)), encoding="utf-8")
            cmd = [sys.executable, str(d / "scripts" / "validate_draft.py"), "--draft", str(fx), "--toc", str(FIX), "--root", str(ROOT)]
        cp = subprocess.run(cmd, capture_output=True, text=True)
        check("T18 설치본 스크립트가 --root로 프로젝트를 찾아 검증 통과", cp.returncode == 0, cp.stdout[-300:] + cp.stderr[-300:])

    # T19
    import lint_project as LP
    with tempfile.TemporaryDirectory() as td:
        t19 = project_copy(td)
        bt = t19 / "book-toc.md"
        p19 = json.loads(json.dumps(load_params(bt)))
        p19["narrator"]["person"] = 2
        p19["review"]["story_pass"] = 95
        p19["characters"].append("")
        p19["anchors"][0]["chapters"].append(9)
        p19["publish"]["ebook"]["epigraph"] = "한 줄"
        write_block(bt, p19)
        _, j = lint_codes(t19)
        ok = all(k in j for k in ("narrator.person", "review.story_pass", "characters[", "없는 장 번호", "publish.ebook.epigraph")) and "book-toc.md:" in j
        check("T19 스키마 범위·빈 항목·장 번호·e북 필드", ok, j[:300])
        txt = bt.read_text(encoding="utf-8").replace('"title": "', '"title" "', 1)
        bt.write_text(txt, encoding="utf-8")
        _, j = lint_codes(t19)
        check("T19 JSON 오류를 한국어로 안내", "JSON 문법 오류" in j and "콜론" in j, j[:200])

    def mk_run(t, run, ledger_text):
        dt, dd = t / "01_test" / run, t / "02_draft" / run
        dt.mkdir(parents=True)
        dd.mkdir(parents=True)
        shutil.copy(t / "book-toc.md", dt / "book-toc.snapshot.md")
        (dd / "00_RUN_STATUS.md").write_text(ledger_text, encoding="utf-8")
        return dt, dd

    # T20
    with tempfile.TemporaryDirectory() as td:
        t20 = project_copy(td)
        mk_run(t20, "20990505_01", "status: in-progress\nwaiver: gate_log 하나\nwaiver: gate_log 둘\ngate: waiver-gate_log → approved \"됨\"\n")
        r, j = lint_codes(t20)
        check("T20 승인 줄 하나는 waiver 하나만", "짝이 되는 사용자 승인 줄이 없다" in j and j.count("waiver 'gate_log'") == 1, j[:300])

    # T21
    with tempfile.TemporaryDirectory() as td:
        t21 = project_copy(td)
        _, dd = mk_run(t21, "20990606_01", "status: in-progress\nstage: storyline done 2099-06-06 (inherited from x)\n")
        (dd / "05_draft-v2.md").write_text("v2", encoding="utf-8")
        (dd / "06_ensemble-1.md").write_text(ensemble(dd / "05_draft-v2.md", 7.0, verdict="재수정", reds=3), encoding="utf-8")
        (dd / "09_draft-final.md").write_text("final", encoding="utf-8")
        _, j = lint_codes(t21)
        check("T21 본문 합평 미통과인데 최종본이 있으면 FAIL", "본문 합평" in j and "통과하지 않았다" in j, j[:300])
        (dd / "06_ensemble-1.md").write_text(ensemble(dd / "05_draft-v2.md", 9.5, verdict="통과", reds=0), encoding="utf-8")
        led = dd / "00_RUN_STATUS.md"
        led.write_text(led.read_text(encoding="utf-8") + "stage: review-5 done 2099-06-06\nreopen: review-9 PROOF 반영\n", encoding="utf-8")
        _, j = lint_codes(t21)
        check("T21 review-9 reopen은 본문 통과를 무르지 않는다", "본문" not in j or "통과 뒤 사용자 승인" not in j, j[:300])

    # T22
    with tempfile.TemporaryDirectory() as td:
        t22 = project_copy(td)
        _, dd = mk_run(t22, "20990707_01", "status: in-progress\n")
        man22 = t22 / "_archive" / "MANIFEST.tsv"
        if not man22.exists():
            man22.write_text("date\tsrc\tdst\treason\n", encoding="utf-8")
        with open(man22, "a", encoding="utf-8") as fh:
            fh.write("2099-07-07\t03_output/20990707_01/ebook.html\t_archive/x/ebook.html\tmoved\n")
        _, j = lint_codes(t22)
        check("T22 승인 없는 작가 산출물 제거 탐지", "[remove_output]" in j, j[:300])
        led = dd / "00_RUN_STATUS.md"
        led.write_text(led.read_text(encoding="utf-8") + 'gate: remove-ebook.html → approved (2099-07-07) "빼도 돼"\n', encoding="utf-8")
        _, j = lint_codes(t22)
        check("T22 승인이 있으면 통과", "remove_output" not in j, j[:300])

    # T23
    with tempfile.TemporaryDirectory() as td:
        t23 = project_copy(td)
        _, dd = mk_run(t23, "20990808_01", "status: published\nstage: storyline done 2099-08-08\nstage: publish done 2099-08-08\n")
        (t23 / "01_test" / "20990808_01" / "storyline.md").write_text("s", encoding="utf-8")
        (dd / "01_research-notes.md").write_text("r", encoding="utf-8")
        (dd / "09_draft-final.md").write_text("f", encoding="utf-8")
        from new_run import open_run
        try:
            open_run(t23, "rework", "다시", "20990808_01", "review", "20991010")
            refused = False
        except SystemExit as e:
            refused = "스토리 게이트" in str(e)
        check("T23 스토리 게이트가 안 닫힌 run은 rework 거부", refused, "")

    # T24
    with tempfile.TemporaryDirectory() as td:
        t24 = project_copy(td)
        shutil.copy(SKILL / "templates" / "storyline.template.md", t24 / "00_user_input" / "storyline.md")
        _, j = lint_codes(t24)
        check("T24 storyline 양식 빈칸 탐지", "양식 빈칸" in j, j[:200])

    # T25
    with tempfile.TemporaryDirectory() as td:
        t25 = project_copy(td)
        dt25, dd = mk_run(t25, "20990909_01", "status: published\nstage: publish done 2099-09-09\n")
        p25 = load_params(dt25 / "book-toc.snapshot.md")
        p25["publish"]["formats"] = ["docx", "ebook"]   # 판형 게이트는 docx를 만드는 run에서만 본다
        write_block(dt25 / "book-toc.snapshot.md", p25)
        _, j = lint_codes(t25)
        check("T25 단계 건너뛰기와 판형 게이트 누락 탐지", "[stage_order]" in j and "판형 게이트" in j, j[:300])

    # T26
    from booktoc import gate_status
    with tempfile.TemporaryDirectory() as td:
        t26 = project_copy(td)
        dt, dd = mk_run(t26, "20991111_01", f"status: in-progress\ninherit: {LIVE}\nstage: storyline done 2099-11-11 (inherited from {LIVE})\nreopen: storyline 다시\n")
        (dt / "storyline_v1.md").write_text("s", encoding="utf-8")
        (dt / "00b_story-ensemble-1.md").write_text(ensemble(dt / "storyline_v1.md", 5.0, verdict="재수정", reds=4), encoding="utf-8")
        ok, why = gate_status("20991111_01", t26, "story")
        check("T26 rework 안의 reopen 뒤 미통과 재합평은 게이트를 연다", not ok, why)
        pp2 = json.loads(json.dumps(load_params(t26 / "book-toc.md")))
        pp2["structure_mode"] = "part"
        pp2["parts"] = [{"no": 1, "title": "가", "chapters": [1, 2]}, {"no": 2, "title": "나", "chapters": [3, 4, 5]}]
        dt2, dd2 = mk_run(t26, "20991111_02", "status: in-progress\n")
        write_block(dt2 / "book-toc.snapshot.md", pp2)
        (dd2 / "05_draft-v2-part1.md").write_text("a", encoding="utf-8")
        (dd2 / "06_ensemble-1-part1.md").write_text(ensemble(dd2 / "05_draft-v2-part1.md", 9.5, verdict="통과"), encoding="utf-8")
        (dd2 / "07_draft-merged-v1.md").write_text("m", encoding="utf-8")
        (dd2 / "06_ensemble-1-overall.md").write_text(ensemble(dd2 / "07_draft-merged-v1.md", 9.5, verdict="통과"), encoding="utf-8")
        ok, why = gate_status("20991111_02", t26, "body")
        check("T26 장편에서 빠진 부(part2)를 잡는다", not ok and "part2" in why, why)
        dt3, dd3 = mk_run(t26, "20991111_03", "status: in-progress\nreopen: review-3 다시\n")
        (dd3 / "05_draft-v2.md").write_text("v2", encoding="utf-8")
        (dd3 / "06_ensemble-1.md").write_text(ensemble(dd3 / "05_draft-v2.md", 9.5, verdict="통과"), encoding="utf-8")
        ok, why = gate_status("20991111_03", t26, "body")
        check("T26 본문 reopen 뒤 끝난 라운드가 없으면 열림", not ok, why)
        dt4, dd4 = mk_run(t26, "20991111_04", 'status: in-progress\nstage: storyline done 2099-11-11\nwaiver: gate_log 본문 합평 누락\ngate: waiver-gate_log → approved (2099-11-11) "됨"\n')
        ok, why = gate_status("20991111_04", t26, "story")
        check("T26 본문용 gate_log waiver는 스토리 게이트를 닫지 않는다", not ok, why)

    # T27
    if LIVE:  # 출판까지 마친 실제 run이 있어야 하는 시험
        with tempfile.TemporaryDirectory() as td:
            t27 = project_copy(td)
            from new_run import open_run
            run, _ = open_run(t27, "rework", "다시 합평", LIVE, "review", "20991212")
            led = (t27 / "02_draft" / run / "00_RUN_STATUS.md").read_text(encoding="utf-8")
            r, j = lint_codes(t27)
            mine = [x for x in r.fails if run in x]   # 원본 run의 상태(다른 세션의 진행 중 수정 등)와 떼어 새 run만 본다
            src_appr = [ln for ln in (ROOT / "02_draft" / LIVE / "00_RUN_STATUS.md").read_text(encoding="utf-8").split("\n")
                        if ln.startswith("gate: waiver-fictional_reviewer")]
            check("T27 rework가 서평 결정을 승인 원문 그대로 물려줌 (새 run FAIL 0)",
                  "waiver: fictional_reviewer" in led and all(s in led for s in src_appr) and not mine, " | ".join(mine)[:300])
            check("T27 rework는 판형을 물려받지 않고(다시 묻는다) 잠금 줄을 남김", "gate: trim →" not in led and "\nlock: " in led, led[-400:])
            fin = (ROOT / "02_draft" / LIVE / "09_draft-final.md").read_text(encoding="utf-8")
            cut = fin.split("\n# 작가의 말")[0]
            f27, _, _ = validate(cut, load_params(ROOT / "01_test" / LIVE / "book-toc.snapshot.md"))
            check("T27 back_matter 섹션을 빼면 FAIL", any("back_matter" in x for x in f27), str(f27)[:200])

    else:
        print("SKIP T27 출판까지 마친 run이 이 저장소에 없다")
    # T28
    with tempfile.TemporaryDirectory() as td:
        t28 = project_copy(td)
        base = 'status: in-progress\nstage: storyline done 2099-01-01\n'
        dt, dd = mk_run(t28, "20990101_01", base)
        led = dd / "00_RUN_STATUS.md"
        led.write_text(base + lock_line("20990101_01", t28) + "\n", encoding="utf-8")
        p28 = load_params(dt / "book-toc.snapshot.md")
        p28["back_matter"] = [x for x in p28["back_matter"] if x != "서평"]
        p28["review"]["story_pass"] = 1.0
        write_block(dt / "book-toc.snapshot.md", p28)
        r, j = lint_codes(t28)
        lk = [x for x in r.fails if "20990101_01 [lock]" in x]
        check("T28 잠금 뒤 서평을 승인 없이 빼면 FAIL", any("서평" in x for x in lk), j[:300])
        check("T28 잠금 뒤 블록 값이 키 승인 없이 바뀌면 FAIL", any("review.story_pass" in x for x in lk), j[:300])
        led.write_text(led.read_text(encoding="utf-8") + 'gate: remove-서평 → approved (2099-01-02) "서평 빼"\ngate: block-change → approved review (2099-01-02) "통과선 바꿔"\n', encoding="utf-8")
        r, j = lint_codes(t28)
        check("T28 승인과 기록이 있으면 잠금 통과", not [x for x in r.fails if "20990101_01 [lock]" in x], j[:300])
        dt2, dd2 = mk_run(t28, "20990101_02", "status: published\nstage: publish done 2099-01-03\n")
        r, j = lint_codes(t28)
        check("T28 출판을 마쳤는데 잠금이 없으면 FAIL", any("20990101_02 [lock]" in x for x in r.fails), j[:300])

    # T29
    from init_project import init as init_project
    with tempfile.TemporaryDirectory() as td:
        t29 = Path(td)
        ans = {"title": "소금의 집", "pen_name": "한별", "genre": "가족 드라마", "genre_label": "단편 소설",
               "target_reader": "부모와 떨어져 사는 30대.", "narrator": {"name": "윤서", "person": 1},
               "characters": ["엄마"], "chapters": [{"title": "부엌", "pages": 14}, {"title": "소금", "pages": 14},
                                                   {"title": "식탁", "pages": 12}], "last_line": "나는 칼을 닦았다."}
        init_project(t29, ans)
        p29 = load_params(t29 / "book-toc.md")
        r, j = lint_codes(t29)
        story_only = [x for x in r.fails if "storyline.md" not in x]
        check("T29 init 블록이 스키마를 통과 (남은 FAIL은 작가가 쓸 storyline 빈칸뿐)",
              p29["length"]["target_pages"] == 40 and p29["chapters"][2]["title"] == "식탁" and not story_only, " | ".join(story_only)[:300])
        st = (t29 / "00_user_input" / "storyline.md").read_text(encoding="utf-8")
        check("T29 storyline에 장 골격과 마지막 줄이 들어감", "### 3장. 식탁 (12매)" in st and "나는 칼을 닦았다." in st, st[:200])
        try:
            init_project(t29, ans)
            again = False
        except SystemExit:
            again = True
        check("T29 있는 파일은 덮어쓰지 않는다", again)

    # T30
    with tempfile.TemporaryDirectory() as td:
        t30 = project_copy(td)
        dt, dd = mk_run(t30, "20990202_01", "")
        p30 = load_params(dt / "book-toc.snapshot.md")
        p30["publish"]["formats"] = ["docx", "ebook"]
        p30["publish"]["ebook"]["epigraph"] = {"lines": ["말"], "attr_lines": ["누군가"], "source_verified": False}
        write_block(dt / "book-toc.snapshot.md", p30)
        out = t30 / "03_output" / "20990202_01"
        out.mkdir(parents=True)
        build(make_draft(load_params(FIX)), p30, "B5", out / "final.docx")
        (dd / "00_RUN_STATUS.md").write_text('status: published\ngate: trim → 신국판 (2099-02-02) "신국판"\nstage: publish done 2099-02-02\n', encoding="utf-8")
        r, j = lint_codes(t30)
        check("T30 원장 판형(신국판)과 다른 docx(B5)를 잡는다", any("20990202_01 [trim]" in x for x in r.fails), j[:300])
        check("T30 승인 없는 미확인 제사를 잡는다", any("20990202_01 [epigraph_source]" in x for x in r.fails), j[:300])
        mk_run(t30, "20990202_02", "status: in-progress\n")
        cp = subprocess.run([sys.executable, str(SKILL / "scripts" / "generate_docx.py"), "--run", "20990202_02", "--root", str(t30)],
                            capture_output=True, text=True, cwd=t30)
        check("T30 판형 결정 없이 docx 생성기가 멈춘다", cp.returncode != 0 and "판형이 정해지지 않았다" in (cp.stderr + cp.stdout), (cp.stderr + cp.stdout)[-300:])

    # T31
    with tempfile.TemporaryDirectory() as td:
        p31 = json.loads(json.dumps(load_params(FIX)))
        p31["publish"]["publisher"] = "테스트출판"
        o = Path(td) / "e.html"
        build_ebook(make_draft(p31), p31, o)
        h = o.read_text(encoding="utf-8")
        check("T31 오프라인 대체 화면 포함", "useOfflineReader" in h and "class SimpleFlip" in h)
        check("T31 디자인 조각이 값과 함께 슬롯으로 옮겨지고 template이 남지 않는다",
              "data-slot=" not in h and h.count("테스트출판") >= 2 and "{{" not in h, h[:0])

    # T32
    with tempfile.TemporaryDirectory() as td:
        t32 = project_copy(td)
        dt, dd = mk_run(t32, "20990303_01", 'status: in-progress\nstage: storyline done 2099-03-03\nwaiver: gate_log story 스토리 게이트 기록 누락\ngate: waiver-gate_log → approved (2099-03-03) "됨"\n')
        (dd / "05_draft-v2.md").write_text("v2", encoding="utf-8")
        (dd / "06_ensemble-1.md").write_text(ensemble(dd / "05_draft-v2.md", 5.0, verdict="재수정", reds=3), encoding="utf-8")
        ok_s, _ = gate_status("20990303_01", t32, "story")
        ok_b, why = gate_status("20990303_01", t32, "body")
        check("T32 스토리용 gate_log waiver는 스토리만 닫고 본문 게이트는 닫지 않는다", ok_s and not ok_b, why)

    # T33
    with tempfile.TemporaryDirectory() as td:
        t33 = project_copy(td)
        base = 'status: in-progress\nstage: storyline done 2099-04-04\n'
        dt, dd = mk_run(t33, "20990404_01", base)
        led = dd / "00_RUN_STATUS.md"
        led.write_text(base + lock_line("20990404_01", t33) + "\n", encoding="utf-8")
        p33 = load_params(dt / "book-toc.snapshot.md")
        p33["back_matter"] = [x for x in p33["back_matter"] if x != "서평"]
        write_block(dt / "book-toc.snapshot.md", p33)
        cp = subprocess.run([sys.executable, str(SKILL / "scripts" / "lock_run.py"), "--run", "20990404_01", "--root", str(t33)],
                            capture_output=True, text=True, cwd=t33)
        check("T33 승인 없는 변경 뒤 lock_run.py가 다시 잠그지 않는다", cp.returncode != 0 and "잠그지 않았다" in cp.stderr + cp.stdout,
              (cp.stderr + cp.stdout)[-200:])
        led.write_text(led.read_text(encoding="utf-8") + lock_line("20990404_01", t33) + "\n", encoding="utf-8")   # 손으로 다시 잠가도
        r, j = lint_codes(t33)
        check("T33 손으로 새 잠금 줄을 써도 앞 구간의 서평 제거를 잡는다",
              any("20990404_01 [lock]" in x and "서평" in x for x in r.fails), j[:300])
        dt2, dd2 = mk_run(t33, "20990404_02", base)
        led2 = dd2 / "00_RUN_STATUS.md"
        led2.write_text(base + lock_line("20990404_02", t33) + "\n" + 'gate: block-change → rejected (2099-04-04) "서평 빼지 마"\n', encoding="utf-8")
        write_block(dt2 / "book-toc.snapshot.md", p33)
        r, j = lint_codes(t33)
        check("T33 rejected 응답은 승인이 아니다", any("20990404_02 [lock]" in x and "서평" in x for x in r.fails), j[:300])
        dt3, dd3 = mk_run(t33, "20990404_03", base)
        led3 = dd3 / "00_RUN_STATUS.md"
        led3.write_text(base + lock_line("20990404_03", t33) + "\nnote: block-change 동기화\n", encoding="utf-8")
        p3 = load_params(dt3 / "book-toc.snapshot.md")
        p3["publish"]["publisher"] = "다른출판사"
        write_block(dt3 / "book-toc.snapshot.md", p3)
        r, j = lint_codes(t33)
        check("T33 note로는 출판 정보를 바꿀 수 없다", any("20990404_03 [lock]" in x and "publish.publisher" in x for x in r.fails), j[:300])

    # T34
    from booktoc import waiver_approvals
    led34 = 'waiver: fictional_reviewer A\nwaiver: fictional_reviewer B\ngate: waiver-fictional_reviewer → approved "A만"\n'
    check("T34 waiver 두 줄 뒤 승인 한 줄은 바로 위(B)만 승인", waiver_approvals(led34) == {1: 2}, str(waiver_approvals(led34)))
    led34b = 'waiver: open_red x\ngate: waiver-open_red → rejected "아니"\ngate: waiver-open_red → approved "응"\n'
    check("T34 사이에 다른 gate가 끼면 짝이 아니다", waiver_approvals(led34b) == {}, str(waiver_approvals(led34b)))

    # T35
    if LIVE:  # 출판까지 마친 실제 run이 있어야 하는 시험
        with tempfile.TemporaryDirectory() as td:
            t35 = project_copy(td)
            run35 = LIVE
            rp35 = [sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", run35, "--root", str(t35)]
            led = t35 / "02_draft" / run35 / "00_RUN_STATUS.md"   # 실제 원장에 남은 post-audit 승인과 떼려고 먼저 잠근다
            led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + "\n" + lock_line(run35, t35) + "\n", encoding="utf-8")
            cp = subprocess.run(rp35, capture_output=True, text=True, cwd=t35)
            check("T35 post-audit 승인 없이 republish.py가 멈춘다", cp.returncode != 0 and "post-audit" in cp.stderr + cp.stdout, (cp.stderr + cp.stdout)[-200:])
            led.write_text(led.read_text(encoding="utf-8") + 'gate: post-audit → approved (2099-05-05) "윤문 반영해"\n', encoding="utf-8")
            cp = subprocess.run(rp35, capture_output=True, text=True, cwd=t35)
            check("T35 --prepare 없이 republish.py가 멈춘다", cp.returncode != 0 and "--prepare" in cp.stderr + cp.stdout, (cp.stderr + cp.stdout)[-200:])
            cp = subprocess.run(rp35 + ["--prepare"], capture_output=True, text=True, cwd=t35)
            mo = re.search(r"owner: (\w+)", cp.stdout)
            fin35 = rp35 + ["--owner", mo.group(1) if mo else "?"]
            arcs = list((t35 / "_archive").glob("*_post-audit*/02_draft/" + run35 + "/09_draft-final.md"))
            check("T35 --prepare가 고치기 전 최종본을 보관하고 MANIFEST에 적는다",
                  cp.returncode == 0 and arcs and "copy before post-audit edit" in (t35 / "_archive" / "MANIFEST.tsv").read_text(encoding="utf-8"),
                  (cp.stderr + cp.stdout)[-200:])
            fin = t35 / "02_draft" / run35 / "09_draft-final.md"
            fin.write_text(fin.read_text(encoding="utf-8").replace("\n\n", "\n\n덧붙인 문장.\n\n", 3), encoding="utf-8")
            cp = subprocess.run(fin35, capture_output=True, text=True, cwd=t35)
            check("T35 변경 기록을 쓰지 않으면 멈춘다", cp.returncode != 0 and "14_post-audit-changelog" in cp.stderr + cp.stdout, (cp.stderr + cp.stdout)[-200:])
            cl = t35 / "02_draft" / run35 / "14_post-audit-changelog.md"
            cl.write_text(cl.read_text(encoding="utf-8") + "\n## 2099-05-05 윤문\n- 문장 세 곳 덧붙임\n", encoding="utf-8")
            bad = fin.read_text(encoding="utf-8")
            fin.write_text(bad.replace("덧붙인 문장.", "덧붙인 — 문장.", 1), encoding="utf-8")
            o35 = next(t35 / "03_output" / run35 / n_ for n_ in ("final.docx", "ebook.html") if (t35 / "03_output" / run35 / n_).exists())
            before = o35.read_bytes()
            cp = subprocess.run(fin35, capture_output=True, text=True, cwd=t35)
            check("T35 원고 검증이 실패하면 산출물을 건드리지 않고 멈춘다",
                  cp.returncode != 0 and o35.read_bytes() == before, (cp.stderr + cp.stdout)[-200:])
            fin.write_text(bad, encoding="utf-8")
            cp = subprocess.run(fin35, capture_output=True, text=True, cwd=t35)
            r, j = lint_codes(t35)
            mine = [x for x in r.fails if run35 in x and x not in BASELINE]
            tail = led.read_text(encoding="utf-8").rstrip().split("\n")[-2:]
            check("T35 승인 뒤 republish.py가 다시 만들고 메타를 맞춰 잠근다 (새 FAIL 0)",
                  cp.returncode == 0 and not mine and tail[0].startswith("stage: post-audit-fix done") and tail[1].startswith("lock:"),
                  ((cp.stderr + cp.stdout)[-200:] + " | " + " | ".join(mine))[:400])

    else:
        print("SKIP T35 출판까지 마친 run이 이 저장소에 없다")
    # T36
    if LIVE:  # 출판까지 마친 실제 run이 있어야 하는 시험
        with tempfile.TemporaryDirectory() as td:
            t36 = project_copy(td)
            from new_run import open_run
            run36, _ = open_run(t36, "rework", "재합평", LIVE, "review", "20991313")
            ae = t36 / "02_draft" / run36 / "00_author-edits.inherited.md"
            check("T36 rework가 출판 뒤 작가 수정 기록을 넘긴다",
                  ae.exists() and "note: author-edits" in (t36 / "02_draft" / run36 / "00_RUN_STATUS.md").read_text(encoding="utf-8"))

    else:
        print("SKIP T36 출판까지 마친 run이 이 저장소에 없다")
    # T37
    import zipfile
    import docx as _docx
    with tempfile.TemporaryDirectory() as td:
        p37 = load_params(FIX)
        o = Path(td) / "a.docx"
        build(make_draft(p37), p37, "신국판", o)
        d = _docx.Document(str(o))
        s0, s1 = d.sections[0], d.sections[-1]
        front_empty = not any(x.text.strip() for x in s0.header.paragraphs + s0.footer.paragraphs) and not s0.footer._element.xpath(".//w:fldChar")
        start = s1._sectPr.xpath("./w:pgNumType/@w:start")
        check("T37 앞부분 구역에는 머리글·쪽번호가 없고 본문 쪽번호는 1부터",
              len(d.sections) == 2 and front_empty and not s1.header.is_linked_to_previous and start == ["1"], str((len(d.sections), front_empty, start)))
        empties = sum(1 for para in d.paragraphs if not para.text.strip())
        check("T37 문서 어디에도 빈 문단이 없다(세로 위치·쪽 나눔·문단 간격은 문단 서식으로)", empties == 0, f"빈 문단 {empties}개")
        ft = zipfile.ZipFile(o).read("word/fontTable.xml").decode("utf-8")
        fb = p37["publish"].get("body_font_fallback") or []
        check("T37 글꼴 표에 계열·한글 문자 집합·대체 이름이 들어간다",
              '<w:family w:val="roman"/>' in ft and '<w:family w:val="swiss"/>' in ft and 'w:charset w:val="81"' in ft and "w:altName" in ft, ft[-300:])

    # T38
    with tempfile.TemporaryDirectory() as td:
        o = Path(td) / "e.html"
        build_ebook(make_draft(load_params(FIX)), load_params(FIX), o)
        h = o.read_text(encoding="utf-8")
        check("T38 200 응답이어도 St.PageFlip이 없으면 다음 후보로", "(window.St && window.St.PageFlip) ? window.dispatchEvent" in h)
        check("T38 글꼴 CSS가 화면을 막지 않고, 대기 상한과 중복 시작 방지가 있다",
              'media="print"' in h and "__fontCss" in h and "__ebookStarted" in h)
        # 행동 시험(Playwright가 있을 때만): 글꼴 CSS와 page-flip이 모두 응답하지 않아도 12초 안에 열린다.
        # 붙잡아 둔 요청을 닫을 때 Playwright가 내는 잡음을 가두려고 따로 띄운 프로세스에서 돌린다
        probe = """
import sys
from playwright.sync_api import sync_playwright
with sync_playwright() as pw:
    br = pw.chromium.launch()
    pg = br.new_context(viewport={"width": 1280, "height": 860}).new_page()
    pg.route("**/*.css", lambda r: None)
    pg.route("**/*page-flip*", lambda r: None)
    pg.goto(sys.argv[1], wait_until="commit")
    try:
        pg.wait_for_function("!document.getElementById('loader') && /\\\\d+ \\\\/ \\\\d+/.test(document.getElementById('counter').textContent)", timeout=12000)
        print("OPENED")
    except Exception:
        print("TIMEOUT")
    br.close()
"""
        cp = subprocess.run([sys.executable, "-c", probe, o.as_uri()], capture_output=True, text=True, timeout=90)
        if "OPENED" in cp.stdout or "TIMEOUT" in cp.stdout:
            check("T38 (브라우저) CSS·라이브러리가 모두 응답하지 않아도 12초 안에 열린다", "OPENED" in cp.stdout)
        else:
            print("SKIP T38 브라우저 시험 (Playwright 없음)")

    # T39
    with tempfile.TemporaryDirectory() as td:
        t39 = project_copy(td)
        base = 'status: in-progress\ngate: new-run → same (2099-06-06) "첫 run"\nstage: storyline done 2099-06-06\n'
        dt, dd = mk_run(t39, "20990606_01", base)
        led = dd / "00_RUN_STATUS.md"
        led.write_text(base + lock_line("20990606_01", t39) + "\n", encoding="utf-8")
        p39 = load_params(dt / "book-toc.snapshot.md")
        p39["publish"]["third_party_verified"] = True
        p39["review"]["body_pass"] = 1.0
        write_block(dt / "book-toc.snapshot.md", p39)
        led.write_text(led.read_text(encoding="utf-8") + 'gate: post-audit → approved (2099-06-06) "오타 하나 고쳐줘"\nnote: block-change review.body_pass\n', encoding="utf-8")
        r, j = lint_codes(t39)
        lk = [x for x in r.fails if "20990606_01 [lock]" in x]
        check("T39 post-audit 승인으로 publish.third_party_verified를 못 바꾼다", any("third_party_verified" in x for x in lk), j[:300])
        check("T39 note로 통과선(review.body_pass)을 못 바꾼다", any("review.body_pass" in x and "note로 바꿀 수 없다" in x for x in lk), j[:300])
        led.write_text(led.read_text(encoding="utf-8") + 'gate: block-change → approved third_party_verified·review (2099-06-06) "실제 서평이고 통과선도 낮춰"\n', encoding="utf-8")
        r, j = lint_codes(t39)
        check("T39 키를 적은 block-change 승인은 통과", not [x for x in r.fails if "20990606_01 [lock]" in x], j[:300])

    # T40
    with tempfile.TemporaryDirectory() as td:
        t40 = project_copy(td)
        base = 'status: in-progress\ngate: new-run → same (2099-07-07) "첫 run"\n'
        dt, dd = mk_run(t40, "20990707_01", base)
        led = dd / "00_RUN_STATUS.md"
        led.write_text(base + lock_line("20990707_01", t40) + "\n", encoding="utf-8")
        led.write_text(base, encoding="utf-8")   # 잠금 줄을 지운다
        r, j = lint_codes(t40)
        check("T40 잠금 줄을 지우면 잡는다", any("20990707_01 [lock]" in x and "지워졌다" in x for x in r.fails), j[:300])

    # T41
    if LIVE:
        with tempfile.TemporaryDirectory() as td:
            t41 = project_copy(td)
            led = t41 / "02_draft" / LIVE / "00_RUN_STATUS.md"
            led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + "\n" + lock_line(LIVE, t41) + "\n", encoding="utf-8")
            fin = t41 / "02_draft" / LIVE / "09_draft-final.md"
            txt = fin.read_text(encoding="utf-8")
            i = txt.index("다.", txt.index("# 1장"))
            fin.write_text(txt[:i] + "요" + txt[i + 1:], encoding="utf-8")   # 한 글자 바꿈, 글자 수 같음
            r, j = lint_codes(t41)
            check("T41 출판본 원고를 승인 없이 한 글자 고치면 잡는다", any(f"{LIVE} [lock]" in x and "최종 원고" in x for x in r.fails), j[:300])
            fin.write_text(txt, encoding="utf-8")
            led.write_text(led.read_text(encoding="utf-8") + 'gate: regenerate-outputs → keep (2099-08-08) "지금 파일 유지"\n', encoding="utf-8")
            cp = subprocess.run([sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", LIVE, "--root", str(t41), "--prepare"],
                                capture_output=True, text=True, cwd=t41)
            check("T41 작가가 유지를 고르면 republish.py가 멈춘다", cp.returncode != 0, (cp.stderr + cp.stdout)[-200:])
    else:
        print("SKIP T41 출판까지 마친 run이 이 저장소에 없다")

    # T42
    if LIVE:
        with tempfile.TemporaryDirectory() as td:
            t42 = project_copy(td)
            led = t42 / "02_draft" / LIVE / "00_RUN_STATUS.md"
            led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + "\n" + lock_line(LIVE, t42) + "\n", encoding="utf-8")
            lf = sorted((t42 / "01_test" / LIVE / "locks").glob("lock-*.json"))[-1]
            orig_lock = lf.read_text(encoding="utf-8")
            lf.write_text(orig_lock.replace('"body_pass": ', '"body_pass": 0.5, "x": ', 1), encoding="utf-8")
            r, j = lint_codes(t42)
            check("T42 잠금 파일만 고치면 잡는다", any(f"{LIVE} [lock]" in x and "해시와 다르다" in x for x in r.fails), j[:300])
            lf.write_text(orig_lock, encoding="utf-8")
            docx_p = next(t42 / "03_output" / LIVE / n_ for n_ in ("final.docx", "ebook.html") if (t42 / "03_output" / LIVE / n_).exists())
            orig_docx = docx_p.read_bytes()
            docx_p.write_bytes(orig_docx + b"\0")
            r, j = lint_codes(t42)
            check("T42 출판 산출물을 승인 없이 바꾸면 잡는다", any(f"{LIVE} [lock]" in x and "출판 산출물" in x for x in r.fails), j[:300])
            docx_p.write_bytes(orig_docx)
            rp = [sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", LIVE, "--root", str(t42)]
            led.write_text(led.read_text(encoding="utf-8") + 'gate: post-audit → approved (2099-09-09) "고쳐"\n', encoding="utf-8")
            c1 = subprocess.run(rp + ["--prepare"], capture_output=True, text=True, cwd=t42)
            c2 = subprocess.run(rp + ["--prepare"], capture_output=True, text=True, cwd=t42)
            check("T42 출판 뒤 수정이 진행 중이면 두 번째 --prepare가 멈춘다", c1.returncode == 0 and c2.returncode != 0 and "끝나지 않았다" in c2.stderr + c2.stdout,
                  (c2.stderr + c2.stdout)[-200:])
            c3 = subprocess.run(rp + ["--owner", "zzzzzz"], capture_output=True, text=True, cwd=t42)
            check("T42 다른 owner는 마칠 수 없다", c3.returncode != 0 and "owner" in c3.stderr + c3.stdout, (c3.stderr + c3.stdout)[-200:])
            from new_run import open_run
            try:
                open_run(t42, "rework", "재합평", LIVE, "review", "20990909")
                blocked = False
            except SystemExit:
                blocked = True
            check("T42 출판 뒤 수정이 진행 중이면 rework가 열리지 않는다", blocked)
        with tempfile.TemporaryDirectory() as td:
            t42b = project_copy(td)
            base = 'status: in-progress\ngate: new-run → same (2099-09-10) "첫 run"\n'
            dt, dd = mk_run(t42b, "20990910_01", base)
            led = dd / "00_RUN_STATUS.md"
            led.write_text(base + lock_line("20990910_01", t42b) + "\n", encoding="utf-8")
            p42 = load_params(dt / "book-toc.snapshot.md")
            p42["review"]["body_pass"] = 1.0
            write_block(dt / "book-toc.snapshot.md", p42)
            led.write_text(led.read_text(encoding="utf-8") + 'gate: block-change → approved publisher (2099-09-10) “출판사 표기만, review 받은 대로”\n', encoding="utf-8")
            r, j = lint_codes(t42b)
            check("T42 응답 원문 속 낱말(review)은 키가 아니다", any("20990910_01 [lock]" in x and "review.body_pass" in x for x in r.fails), j[:300])
    else:
        print("SKIP T42 출판까지 마친 run이 이 저장소에 없다")

    # T43
    if LIVE:
        with tempfile.TemporaryDirectory() as td:
            t43 = project_copy(td)
            led = t43 / "02_draft" / LIVE / "00_RUN_STATUS.md"
            led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + "\n" + lock_line(LIVE, t43) + "\n", encoding="utf-8")
            fin = t43 / "02_draft" / LIVE / "09_draft-final.md"
            fin.write_text(fin.read_text(encoding="utf-8").replace("다.", "요.", 1), encoding="utf-8")
            led.write_text(led.read_text(encoding="utf-8") + 'gate: post-audit → approved (2099-10-10) "고쳐"\n', encoding="utf-8")
            cp = subprocess.run([sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", LIVE, "--root", str(t43), "--prepare"],
                                capture_output=True, text=True, cwd=t43)
            check("T43 승인 전 수정이 있으면 --prepare가 멈춘다", cp.returncode != 0 and "이미 바뀌었다" in cp.stderr + cp.stdout, (cp.stderr + cp.stdout)[-200:])
    else:
        print("SKIP T43 출판까지 마친 run이 이 저장소에 없다")

    # T44
    import generate_docx as _gd
    sample = next((p for d in ("/System/Library/Fonts", "C:/Windows/Fonts", "/usr/share/fonts")
                   for p in (Path(d).rglob("*.tt[fc]") if Path(d).exists() else [])), None)
    if sample:
        check("T44 글꼴 파일의 name 표에서 계열 이름을 읽는다", bool(_gd.font_families(sample)), str(sample))
    else:
        print("SKIP T44 시스템 글꼴 폴더를 찾지 못했다")

    # T45
    if LIVE:
        with tempfile.TemporaryDirectory() as td:
            t45 = project_copy(td)
            led = t45 / "02_draft" / LIVE / "00_RUN_STATUS.md"
            led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + "\n" + lock_line(LIVE, t45) + "\n"
                           + 'gate: post-audit → approved (2099-11-11) "고쳐"\n', encoding="utf-8")
            rp = [sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", LIVE, "--root", str(t45)]
            run_ = lambda extra: subprocess.run(rp + extra, capture_output=True, text=True, cwd=t45)  # noqa: E731
            run_(["--prepare"])
            fin = t45 / "02_draft" / LIVE / "09_draft-final.md"
            orig = (sorted((t45 / "_archive").glob(f"*_post-audit_*/02_draft/{LIVE}/09_draft-final.md"))[-1]).read_text(encoding="utf-8")
            fin.write_text(orig.replace("다.", "요.", 1), encoding="utf-8")   # 세션이 반쯤 고치다 끊김
            st = run_(["--status"])
            check("T45 --status가 진행 중인 수정을 알린다", "진행 중인 출판 뒤 수정: 원장" in st.stdout, st.stdout)
            c1 = run_(["--abandon"])
            check("T45 승인 없이는 버리지 않는다", c1.returncode != 0 and "post-audit-abandon" in c1.stderr + c1.stdout, (c1.stderr + c1.stdout)[-200:])
            led.write_text(led.read_text(encoding="utf-8") + 'gate: post-audit-abandon → approved (2099-11-11) "그 수정은 버려"\n', encoding="utf-8")
            c2 = run_(["--abandon"])
            r, j = lint_codes(t45)
            check("T45 버리기는 원고를 수정 전으로 되돌리고 수정을 닫는다",
                  c2.returncode == 0 and fin.read_text(encoding="utf-8") == orig and "진행 중인 출판 뒤 수정 없음" in run_(["--status"]).stdout
                  and not [y for y in r.fails if f"{LIVE} [lock]" in y], ((c2.stderr + c2.stdout)[-200:] + j[:200]))
            led.write_text(led.read_text(encoding="utf-8") + 'gate: post-audit → approved (2099-11-12) "다시 고쳐"\n', encoding="utf-8")
            run_(["--prepare"])
            led.write_text(led.read_text(encoding="utf-8") + 'gate: post-audit-takeover → approved (2099-11-12) "이어받아"\n', encoding="utf-8")
            c3 = run_(["--takeover"])
            mo = re.search(r"새 owner: (\w+)", c3.stdout)
            cl = t45 / "02_draft" / LIVE / "14_post-audit-changelog.md"
            cl.write_text(cl.read_text(encoding="utf-8") + "\n## 2099-11-12 이어받아 마침\n", encoding="utf-8")
            c4 = run_(["--owner", mo.group(1) if mo else "?"])
            check("T45 이어받기는 새 owner로 마친다", c3.returncode == 0 and c4.returncode == 0, (c3.stdout + c4.stderr + c4.stdout)[-300:])
            r, j = lint_codes(t45)
            check("T45 버리기 → 새 수정 → 이어받아 마친 뒤 린터에 새 잠금 FAIL이 없다",
                  not [y for y in r.fails if f"{LIVE} [lock]" in y or f"{LIVE} [post_audit]" in y], j[:300])
            c5 = run_(["--takeover"])
            check("T45 이어받기 승인 하나로 두 번 이어받지 못한다", c5.returncode != 0, (c5.stderr + c5.stdout)[-200:])
    else:
        print("SKIP T45 출판까지 마친 run이 이 저장소에 없다")

    # T46
    if LIVE:
        with tempfile.TemporaryDirectory() as td:
            t46 = project_copy(td)
            led = t46 / "02_draft" / LIVE / "00_RUN_STATUS.md"
            led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + "\n" + lock_line(LIVE, t46) + "\n"
                           + 'gate: post-audit → approved (2099-12-01) "오타 하나"\n', encoding="utf-8")
            rp = [sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", LIVE, "--root", str(t46)]
            fin = t46 / "02_draft" / LIVE / "09_draft-final.md"
            orig = fin.read_text(encoding="utf-8")
            fin.write_text(orig.replace("다.", "요.", 1), encoding="utf-8")
            led.write_text(led.read_text(encoding="utf-8") + "stage: post-audit-fix done 2099-12-01 (1장 오타)\n", encoding="utf-8")
            cp = subprocess.run([sys.executable, str(SKILL / "scripts" / "lock_run.py"), "--run", LIVE, "--root", str(t46)],
                                capture_output=True, text=True, cwd=t46)
            r, j = lint_codes(t46)
            check("T46 손으로 쓴 done 줄로는 닫히지 않는다(재잠금 거부, prepare 없는 수정 FAIL)",
                  cp.returncode != 0 and any(f"{LIVE} [lock]" in y and "--prepare 기록이 없다" in y for y in r.fails), (cp.stderr + cp.stdout)[-150:] + j[:200])
            fin.write_text(orig, encoding="utf-8")
        with tempfile.TemporaryDirectory() as td:
            t46b = project_copy(td)
            led = t46b / "02_draft" / LIVE / "00_RUN_STATUS.md"
            led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + "\n" + lock_line(LIVE, t46b) + "\n"
                           + 'gate: post-audit → approved (2099-12-02) "고쳐"\n', encoding="utf-8")
            rp = [sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", LIVE, "--root", str(t46b)]
            subprocess.run(rp + ["--prepare"], capture_output=True, text=True, cwd=t46b)
            led.write_text(led.read_text(encoding="utf-8") + "note: post-audit-prepare _archive/x\n", encoding="utf-8")   # 다른 세션의 위조 prepare
            fin = t46b / "02_draft" / LIVE / "09_draft-final.md"
            fin.write_text(fin.read_text(encoding="utf-8").replace("다.", "요.", 1), encoding="utf-8")
            cp = subprocess.run(rp, capture_output=True, text=True, cwd=t46b)
            r, j = lint_codes(t46b)
            check("T46 owner 없는 prepare 줄을 덧붙여 마쳐도 잡는다",
                  cp.returncode != 0 or any(f"{LIVE} [lock]" in y and "정확히 하나" in y for y in r.fails), (cp.stderr + cp.stdout)[-150:] + j[:200])
    else:
        print("SKIP T46 출판까지 마친 run이 이 저장소에 없다")

    # T47
    from booktoc import leaf_allowed
    old = {"publish": {"publisher": "<출판사>", "pub_date": "2026년 5월",
                       "ebook": {"epigraph": {"lines": ["a"], "source_verified": False}, "back_quote": {"lines": ["b"]}}}}
    check("T47 겹치는 끝 이름(lines)은 다른 경로를 덮지 않는다",
          not leaf_allowed("publish.ebook.back_quote.lines", 'gate: block-change → approved lines (d) "x"', {"block-change"}, old))
    check("T47 publish-info는 빈칸을 채운 것만 덮는다",
          leaf_allowed("publish.publisher", 'gate: publish-info → approved (d) "x"', {"publish-info"}, old)
          and not leaf_allowed("publish.pub_date", 'gate: publish-info → approved (d) "x"', {"publish-info"}, old))
    check("T47 epigraph-source는 출처 표시만 덮는다",
          leaf_allowed("publish.ebook.epigraph.source_verified", 'gate: epigraph-source → verified (d) "x"', {"epigraph-source"}, old)
          and not leaf_allowed("publish.ebook.epigraph.lines", 'gate: epigraph-source → verified (d) "x"', {"epigraph-source"}, old))

    # T48
    if LIVE:
        with tempfile.TemporaryDirectory() as td:
            t48 = project_copy(td)
            meta = t48 / "03_output" / LIVE / "metadata.md"
            meta.write_text(re.sub(r"분량:\s*[\d,]+자", "분량: 약 1만8천 자", meta.read_text(encoding="utf-8")), encoding="utf-8")
            r, j = lint_codes(t48)
            check("T48 metadata 분량 줄 형식이 다르면 출판 감사가 잡는다", any(f"{LIVE} [metadata]" in y and "형식" in y for y in r.fails), j[:300])
    msgs = []
    for sc in ("validate_draft.py", "generate_docx.py", "generate_ebook.py", "lock_run.py"):
        cp = subprocess.run([sys.executable, str(SKILL / "scripts" / sc), "--run", "20991231_99", "--root", str(ROOT)], capture_output=True, text=True)
        msgs.append(cp.returncode != 0 and "있는 run" in cp.stderr + cp.stdout and "Traceback" not in cp.stderr)
    check("T48 없는 run이면 네 스크립트 모두 traceback 없이 있는 run 목록으로 답한다", all(msgs), str(msgs))

    # T49
    with tempfile.TemporaryDirectory() as td:
        p49 = json.loads(json.dumps(load_params(FIX)))
        p49["structure_mode"] = "part"
        p49["parts"] = [{"no": 1, "title": "가", "chapters": [1, 2]}, {"no": 2, "title": "나", "chapters": [3]}]
        o = Path(td) / "p.docx"
        build(make_draft(p49), p49, "신국판", o)
        d = _docx.Document(str(o))
        heads = [bool(s.header.paragraphs[0].text.strip()) for s in d.sections]
        starts = [s._sectPr.xpath("./w:pgNumType/@w:start") for s in d.sections]
        check("T49 부 제목 쪽은 머리글 없음, 장 쪽은 있음, 번호는 한 번만 1로",
              heads == [False, False, True, False, True] and starts == [[], ["1"], [], [], []] and
              not any(not q.text.strip() for q in d.paragraphs), str((heads, starts)))

    # T50 (docx를 만드는 run이 아니면 사본에서 docx를 만드는 run으로 바꿔 시험한다)
    if LIVE:
        with tempfile.TemporaryDirectory() as td:
            t50 = project_copy(td)
            snap = t50 / "01_test" / LIVE / "book-toc.snapshot.md"
            p50 = load_params(snap)
            if "docx" not in p50["publish"].get("formats", []):
                p50["publish"]["formats"] = ["docx"] + p50["publish"].get("formats", [])
                write_block(snap, p50)
            led = t50 / "02_draft" / LIVE / "00_RUN_STATUS.md"
            if "gate: trim →" not in led.read_text(encoding="utf-8"):
                led.write_text(led.read_text(encoding="utf-8") + 'gate: trim → 신국판 (selftest) "신국판"\n', encoding="utf-8")
            dx = t50 / "03_output" / LIVE / "final.docx"
            fin = (t50 / "02_draft" / LIVE / "09_draft-final.md").read_text(encoding="utf-8")
            import contextlib
            import io
            with contextlib.redirect_stdout(io.StringIO()):
                build(fin, p50, "신국판", dx)
            d = _docx.Document(str(dx))
            d.paragraphs[-1].insert_paragraph_before("")   # 옛 생성기처럼 빈 문단이 낀 docx
            d.save(str(dx))
            r = lint(t50)
            check("T50 출판된 docx가 지금 생성기 결과와 다르면 WARN", any("final.docx가 지금 원고·생성기로" in w for w in r.warns), " | ".join(r.warns)[:300])
    else:
        print("SKIP T50 출판까지 마친 run이 이 저장소에 없다")

    # T51
    if LIVE:
        with tempfile.TemporaryDirectory() as td:
            t51 = project_copy(td)
            led = t51 / "02_draft" / LIVE / "00_RUN_STATUS.md"
            led.write_text(led.read_text(encoding="utf-8").rstrip("\n") + "\n" + lock_line(LIVE, t51) + "\n"
                           + 'gate: post-audit → approved (2099-12-20) "서평 문장 고쳐"\n', encoding="utf-8")
            rp = [sys.executable, str(SKILL / "scripts" / "republish.py"), "--run", LIVE, "--root", str(t51)]
            c0 = subprocess.run(rp + ["--prepare"], capture_output=True, text=True, cwd=t51)
            owner = re.search(r"owner: (\w+)", c0.stdout).group(1)
            out = t51 / "03_output" / LIVE
            keep = t51 / "_archive" / f"20991220_000000_000000_post-audit_{owner}" / "03_output" / LIVE   # 마침 단계가 남긴 보관본처럼
            keep.mkdir(parents=True)
            for n_ in ("final.docx", "ebook.html", "metadata.md"):
                if (out / n_).exists():
                    shutil.copy2(out / n_, keep / n_)
            fin = t51 / "02_draft" / LIVE / "09_draft-final.md"
            fin.write_text(fin.read_text(encoding="utf-8").replace("다.", "요.", 1), encoding="utf-8")
            (out / "ebook.html").write_text((out / "ebook.html").read_text(encoding="utf-8") + "<!-- 버려질 문장 -->", encoding="utf-8")
            st = subprocess.run(rp + ["--status"], capture_output=True, text=True, cwd=t51).stdout
            check("T51 --status가 요청 원문과 원고·산출물 변경 여부를 보여 준다",
                  "서평 문장 고쳐" in st and "원고가 잠금 뒤 바뀌었나: 예" in st and "docx·e북이 잠금 뒤 바뀌었나: 예" in st, st[:300])
            led.write_text(led.read_text(encoding="utf-8") + 'gate: post-audit-abandon → approved (2099-12-21) "버려"\n', encoding="utf-8")
            c1 = subprocess.run(rp + ["--abandon"], capture_output=True, text=True, cwd=t51)
            r, j = lint_codes(t51)
            check("T51 버리기가 산출물까지 되돌린다", c1.returncode == 0 and "버려질 문장" not in (out / "ebook.html").read_text(encoding="utf-8")
                  and not [y for y in r.fails if f"{LIVE} [lock]" in y], (c1.stderr + c1.stdout)[-200:] + j[:200])
    else:
        print("SKIP T51 출판까지 마친 run이 이 저장소에 없다")

    n_fail = results.count(False)
    print(f"\nselftest: {len(results) - n_fail}/{len(results)} PASS")
    sys.exit(1 if n_fail else 0)


if __name__ == "__main__":
    main()
